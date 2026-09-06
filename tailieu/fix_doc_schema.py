import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__))).replace('\\', '/')

def fix_document_schema(doc_path):
    print(f"Fixing schema in {doc_path}...")
    doc = docx.Document(doc_path)
    
    # Text adjustments: 6 TUẦN
    for p in doc.paragraphs:
        if "9 TUẦN" in p.text:
            p.text = p.text.replace("9 TUẦN", "6 TUẦN")
        if "9 tuần" in p.text:
            p.text = p.text.replace("9 tuần", "6 tuần")
            
    plan_data = [
        # W1
        ("Tuần 01 (27/07/2026 – 02/08/2026)", "Khảo sát bài toán quản lý bán hàng tại quầy (POS), quản lý kho, nhập hàng, quản trị danh mục sản phẩm và nhu cầu tích hợp AI.", "Nguyễn Viết Cường", "DEL-001: Tóm tắt bài toán và phạm vi dự án."),
        ("", "Thu thập thông tin nhóm, phân chia vai trò (Backend / Frontend / AI / QA), xác định stakeholder và mục tiêu đề tài.", "Hoàng Trang Hiên", "STK-001, OBJ-001; Thống nhất phân công 2 thành viên."),
        ("", "Lập kế hoạch thực hiện dự án 6 tuần theo chuẩn quy trình SDLC, xác định các mốc nộp bài và tiêu chí đánh giá.", "Hoàng Trang Hiên", "MS-001; DEL-002: Project Plan 6 tuần hoàn chỉnh."),
        ("", "Lập kế hoạch sử dụng OpenAI API, chuẩn bị môi trường phát triển (Python 3.12, Node.js, Git) và cấu trúc thư mục dự án.", "Cả nhóm", "Môi trường làm việc chuẩn bị sẵn sàng."),
        
        # W2
        ("Tuần 02 (03/08/2026 – 09/08/2026)", "Làm rõ yêu cầu nghiệp vụ: Bán hàng tại quầy POS, quét mã vạch, tự động trừ tồn kho, quản lý đơn hàng và phân quyền 2 Actor (Admin, Owner).", "Nguyễn Viết Cường", "Chuẩn bị đầu vào cho Requirements QA."),
        ("", "Xây dựng Bảng hỏi làm rõ yêu cầu (Requirements QA: QA-001..QA-020) và xác định các quy tắc nghiệp vụ Business Rules (BR-001..BR-004).", "Hoàng Trang Hiên", "DEL-003: Requirements QA làm rõ phạm vi nghiệp vụ."),
        ("", "Đặc tả yêu cầu Module AI: Chatbot tư vấn chọn sản phẩm (Consultant) và Phân tích doanh thu kinh doanh (Revenue Insights).", "Hoàng Trang Hiên", "REQ-AI; Thiết lập Guardrails bảo vệ dữ liệu nhạy cảm."),
        ("", "Hoàn thiện tài liệu Đặc tả Yêu cầu Phần mềm (SRS) và ma trận truy vết yêu cầu (Traceability Matrix).", "Cả nhóm", "DEL-004: SRS phục vụ thiết kế kiến trúc hệ thống."),

        # W3
        ("Tuần 03 (10/08/2026 – 16/08/2026)", "Thiết kế Sơ đồ Use Case tổng quát chuẩn UML 2.5 với 2 Actor chính (Admin, Owner) và 10 Use Cases cốt lõi kèm quan hệ «include».", "Nguyễn Viết Cường", "Diagram 01: Sơ đồ Use Case Tổng quát."),
        ("", "Thiết kế Sơ đồ Lớp Lĩnh vực (Domain Class Diagram) với 7 Entity Classes + AIService, quan hệ Composition, Multiplicity và Phương thức.", "Hoàng Trang Hiên", "Diagram 02: Sơ đồ Lớp Lĩnh vực (Domain Class)."),
        ("", "Thiết kế Kiến trúc hệ thống 3 tầng (3-Tier Architecture: React SPA ↔ FastAPI REST API ↔ SQLite & OpenAI Cloud).", "Nguyễn Viết Cường", "Diagram 04: Sơ đồ Kiến trúc Hệ thống."),
        ("", "Hoàn thiện tài liệu Thiết kế Hướng đối tượng (OOD) và chuẩn bị hồ sơ minh chứng Bài kiểm tra số 1 (KT1).", "Cả nhóm", "DEL-005: Tài liệu OOD; Mốc nộp Bài KT1."),

        # W4
        ("Tuần 04 (17/08/2026 – 23/08/2026)", "Thiết kế Cơ sở dữ liệu quan hệ (ERD 7 bảng cốt lõi trong sales.db: users, categories, products, orders, order_details, customers, receipts).", "Hoàng Trang Hiên", "Diagram 03: Sơ đồ Thực thể CSDL (ERD)."),
        ("", "Thiết kế Sơ đồ Hoạt động (Activity Diagram) cho luồng Bán hàng POS và Luồng Trợ lý AI có Data Masking & Fallback.", "Nguyễn Viết Cường", "Diagram 05 & 06: Sơ đồ Hoạt động POS & AI."),
        ("", "Thiết kế Sơ đồ Tuần tự (Sequence Diagram) cho quy trình giao dịch POS Transaction Commit và trừ tồn kho tự động.", "Hoàng Trang Hiên", "Diagram 07: Sơ đồ Tuần tự Giao dịch POS."),
        ("", "Xây dựng Từ điển dữ liệu (Data Dictionary), chuẩn hóa các ràng buộc toàn vẹn khóa chính PK, khóa ngoại FK và Unique Index.", "Cả nhóm", "DEL-006: Tài liệu Thiết kế CSDL & Flow Design."),

        # W5
        ("Tuần 05 (24/08/2026 – 30/08/2026)", "Lập trình Backend REST API (FastAPI, Python 3.12, SQLAlchemy ORM, JWT Authentication, Quản lý kho, POS, Hóa đơn).", "Nguyễn Viết Cường", "Hoàn thiện các API Endpoints cốt lõi."),
        ("", "Lập trình Frontend POS (React 18 SPA, Vite, Tailwind CSS, Lucide Icons, Giao diện POS bán hàng, Dashboard thống kê, In bill).", "Nguyễn Viết Cường", "Giao diện người dùng tương tác trực quan."),
        ("", "Tích hợp Module OpenAI GPT-4o-mini: Prompt Engineering, Data Masking (ẩn giá vốn/SĐT), Rate Limiting và Fallback an toàn.", "Hoàng Trang Hiên", "DEL-007: Chức năng Bán hàng & Trợ lý AI."),
        ("", "Rà soát kết nối tích hợp hệ thống End-to-End giữa Frontend React, Backend FastAPI và CSDL SQLite.", "Cả nhóm", "Hoàn thành phiên bản hệ thống V1.0."),

        # W6
        ("Tuần 06 (31/08/2026 – 06/09/2026)", "Xây dựng Test Scenario và Test Cases cho bán hàng POS, trừ tồn kho, quản lý danh mục và phân quyền Module AI.", "Hoàng Trang Hiên", "TC-001..TC-020; DEL-008: Báo cáo Test & Bug Log."),
        ("", "Thực hiện kiểm thử chức năng (Functional Testing), sửa lỗi phát sinh và tối ưu hiệu năng tốc độ xử lý đơn hàng.", "Hoàng Trang Hiên", "Hệ thống vận hành ổn định, không lỗi nghiêm trọng."),
        ("", "Viết tài liệu README, Hướng dẫn cài đặt / triển khai và Sổ tay hướng dẫn sử dụng (User Guide) theo phân quyền từng Actor.", "Nguyễn Viết Cường", "DEL-009: README & User Guide."),
        ("", "Đóng gói toàn bộ sản phẩm phần mềm, chuẩn bị Slide báo cáo, kịch bản Demo và nghiệm thu dự án cuối kỳ.", "Cả nhóm", "DEL-010: Hồ sơ nghiệm thu dự án hoàn chỉnh.")
    ]

    t1 = doc.tables[1]
    
    # Remove existing rows except header
    while len(t1.rows) > 1:
        tr = t1.rows[-1]._tr
        tr.getparent().remove(tr)
        
    for r_idx, (w_time, task, member, note) in enumerate(plan_data):
        row = t1.add_row()
        cells = row.cells
        
        bg_color = "EFF6FF" if w_time.startswith("Tuần") else ("F8FAFC" if (r_idx % 2 == 1) else "FFFFFF")
        
        for c_idx, cell in enumerate(cells):
            tcPr = cell._tc.get_or_add_tcPr()
            # Clean old tcMar & shd
            for child in list(tcPr):
                if child.tag.endswith(('shd', 'tcMar')):
                    tcPr.remove(child)
                    
            # Strict schema-compliant shd
            shd_xml = parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{bg_color}"/>')
            tcPr.append(shd_xml)
            
            # Values
            val = [w_time, task, member, note][c_idx]
            cell.text = val
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            for run in p.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(9.5)
                if c_idx == 0 and w_time:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(30, 58, 138)
                elif c_idx == 2 and member == "Cả nhóm":
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(5, 150, 105)
                elif c_idx == 3:
                    run.font.size = Pt(9.0)
                    run.font.color.rgb = RGBColor(71, 85, 105)

    doc.save(doc_path)
    print(f"Fixed {doc_path} successfully!")

if __name__ == "__main__":
    paths = [
        f"{BASE_DIR}/tailieu/BAO_CAO_KT1_NHOM03.docx",
        f"{BASE_DIR}/tailieu/BAO_CAO_KT1_NHOM03_MOI.docx",
        f"{BASE_DIR}/docs/BaoCao_TongHop_QuanLyBanHang_AI_Nhom03.docx"
    ]
    for p in paths:
        try:
            fix_document_schema(p)
        except Exception as e:
            print(f"Error {p}: {e}")